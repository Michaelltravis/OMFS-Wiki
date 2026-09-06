from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader


OUTPUT_DIR = Path(r"C:\Users\micha\Desktop\Richmond\Draft Sections\01 Working Drafts")
PDF_DIR = Path(r"C:\Users\micha\Desktop\Wiki\work\richmond_qa_20260905\pdf")
REPORT_PATH = Path(r"C:\Users\micha\Desktop\Wiki\work\richmond_qa_20260905\final_validation.json")

EXPECTED = [
    "01_Tech_1_Cover_Letter_Richmond_Draft.docx",
    "02_Tech_2_Executive_Summary_Richmond_Draft.docx",
    "03_Tech_3_Firm_Qualifications_Richmond_Draft.docx",
    "04_Tech_4_Staffing_Management_Richmond_Draft.docx",
    "05_Tech_5_Technical_Approach_Richmond_Draft.docx",
    "06_Tech_6_Required_Forms_Richmond_Draft.docx",
    "07_Fee_Proposal_Richmond_Draft.docx",
]
REQUIRED_SNIPPETS = {
    "01_Tech_1_Cover_Letter_Richmond_Draft.docx": [
        "[CONFIRM RECIPIENT ADDRESS]",
        "more than four decades of contract O&M experience",
    ],
    "03_Tech_3_Firm_Qualifications_Richmond_Draft.docx": [
        "10 OR MORE YEARS PROVIDING WASTEWATER-TREATMENT O&M SERVICES",
        "AT LEAST FIVE ACTIVATED-SLUDGE WWTPS OF 10 MGD OR GREATER",
        "AT LEAST FIVE COLLECTION SYSTEMS OF 50 MILES OR GREATER",
        "ANNUAL O&M REVENUE EXCEEDS $50 MILLION",
    ],
    "05_Tech_5_Technical_Approach_Richmond_Draft.docx": [
        "December 31",
        "June 1",
        "June 30",
        "Start within two hours",
        "10th of each month",
        "Self-performance and subcontracting controls",
        "then annually within 30 days after the renewal date",
    ],
    "07_Fee_Proposal_Richmond_Draft.docx": [
        "Work exactly $5,000",
        "Contract Operator-owned equipment",
        "Operator-caused repairs",
        "RFP-defined disposal responsibilities",
        "sodium bisulfate",
        "sodium bisulfite",
    ],
}

TRIP_PREFIX = "Source: Due Diligence Trip Report – "
CLASSIFICATIONS = (
    "Observed condition",
    "Reported condition",
    "Recommendation",
    "Open question/data gap",
)
REQUIRED_STYLE_SIZES = {
    "Normal": 12.0,
    "Source Note": 12.0,
    "Placeholder": 12.0,
    "Table Text": 12.0,
}
ALLOWED_TRIP_REPORT_AUTHORS = {
    "Kelly Irving",
    "Daniel Gorka",
    "Ryan Larson",
    "Roy Aristizabal",
    "Daniel Buonadonna",
    "Trey Cain",
    "Iouri Ossokine",
}


def all_text(doc: Document) -> str:
    parts: list[str] = []
    parts.extend(p.text for p in doc.paragraphs if p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs if p.text)
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs if p.text)
        parts.extend(p.text for p in section.footer.paragraphs if p.text)
    return "\n".join(parts)


def validate_one(path: Path) -> dict:
    package_ok = True
    package_error = None
    try:
        with zipfile.ZipFile(path) as package:
            bad_member = package.testzip()
            if bad_member:
                package_ok = False
                package_error = f"CRC failure: {bad_member}"
            required = {"[Content_Types].xml", "word/document.xml"}
            missing = sorted(required.difference(package.namelist()))
            if missing:
                package_ok = False
                package_error = f"Missing package members: {missing}"
    except Exception as exc:  # pragma: no cover - validation failure path
        package_ok = False
        package_error = str(exc)

    doc = Document(path)
    text = all_text(doc)
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    trip_source_lines = [
        line for line in lines
        if line.startswith("Source:") and "Due Diligence Trip Report" in line
    ]
    malformed_trip_lines = [line for line in trip_source_lines if not line.startswith(TRIP_PREFIX)]
    cited_authors = {
        match.group(1).strip()
        for line in trip_source_lines
        if (match := re.match(re.escape(TRIP_PREFIX) + r"([^,]+),", line))
    }
    register_line = next(
        (line for line in lines if line.startswith("Due Diligence Trip Reports by ")),
        None,
    )
    registered_authors: set[str] = set()
    if register_line:
        names = register_line.removeprefix("Due Diligence Trip Reports by ")
        names = re.sub(r",? cited above\.?$", "", names)
        names = names.rstrip(".").replace(", and ", ", ").replace(" and ", ", ")
        registered_authors = {name.strip() for name in names.split(",") if name.strip()}
    class_counts = {label: len(re.findall(re.escape(label), text, flags=re.I)) for label in CLASSIFICATIONS}
    placeholders = sorted(set(re.findall(r"\[[A-Z][A-Z0-9 /&'().,_-]{1,80}\]", text)))

    pdf_path = PDF_DIR / f"{path.stem}.pdf"
    page_count = len(PdfReader(pdf_path).pages) if pdf_path.exists() else None

    errors: list[str] = []
    warnings: list[str] = []
    if not package_ok:
        errors.append(package_error or "Invalid package")
    if "Draft Source and Decision Register" not in text:
        errors.append("Missing Draft Source and Decision Register")
    for snippet in REQUIRED_SNIPPETS.get(path.name, []):
        if snippet not in text:
            errors.append(f"Missing required draft control text: {snippet}")
    if malformed_trip_lines:
        errors.append(f"Malformed trip-report source notes: {len(malformed_trip_lines)}")
    for line in trip_source_lines:
        labels = re.findall(
            r"Evidence type: (Observed condition|Reported condition|Recommendation|Open question/data gap)\.",
            line,
        )
        if len(labels) != 1:
            errors.append(f"Trip-report source note lacks exactly one evidence label: {line}")
    if any("mack" in line.lower() for line in trip_source_lines):
        errors.append("Excluded Mack template/report appears in trip-report citation")
    unknown_authors = sorted(cited_authors.difference(ALLOWED_TRIP_REPORT_AUTHORS))
    if unknown_authors:
        errors.append(f"Unknown trip-report author(s): {unknown_authors}")
    if cited_authors and not register_line:
        errors.append("Trip-report source notes present without a source-register entry")
    if register_line and registered_authors != cited_authors:
        errors.append(
            "Trip-report source-register mismatch: "
            f"cited={sorted(cited_authors)}, registered={sorted(registered_authors)}"
        )
    if re.search(r"29\s*[–-]\s*32\s+FTE", text, flags=re.I):
        errors.append("Unsupported 29–32 FTE range appears")
    prohibited = {
        "450 Civic Center Plaza": "Unsupported recipient address appears",
        "conflicting questions deadline": "Unsupported deadline-conflict statement appears",
        "on-site corrective-maintenance labor": "Corrective-maintenance labor is improperly narrowed to on-site labor",
        "Future bioswales, bioretention, and low-impact-development assets are excluded": "Unsupported future-asset exclusion appears",
        "Return to City or roll over at City option": "Unsupported M&R rollover option appears",
    }
    for phrase, message in prohibited.items():
        if phrase.lower() in text.lower():
            errors.append(message)
    style_sizes = {}
    for style_name, required_size in REQUIRED_STYLE_SIZES.items():
        style = doc.styles[style_name]
        actual_size = style.font.size.pt if style.font.size else None
        style_sizes[style_name] = actual_size
        if actual_size is None or abs(actual_size - required_size) > 0.01:
            errors.append(f"{style_name} style is {actual_size} pt; expected {required_size} pt")
    if trip_source_lines and not any(class_counts.values()):
        errors.append("Trip-report evidence is present without an evidence classification")
    if page_count is None:
        errors.append("Missing native-Word QA PDF")
    if not trip_source_lines and path.name not in {
        "06_Tech_6_Required_Forms_Richmond_Draft.docx",
        "07_Fee_Proposal_Richmond_Draft.docx",
    }:
        warnings.append("No trip-report citation found")

    return {
        "file": path.name,
        "bytes": path.stat().st_size,
        "package_ok": package_ok,
        "pages": page_count,
        "words": len(re.findall(r"\b[\w’'-]+\b", text)),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "trip_report_source_notes": len(trip_source_lines),
        "trip_report_authors": sorted(cited_authors),
        "classification_mentions": class_counts,
        "placeholders": placeholders,
        "required_style_sizes_pt": style_sizes,
        "errors": errors,
        "warnings": warnings,
    }


def main() -> int:
    actual = sorted(p.name for p in OUTPUT_DIR.glob("*.docx") if not p.name.startswith("~$"))
    top_errors: list[str] = []
    if actual != EXPECTED:
        top_errors.append(f"Expected exactly seven named DOCX files; found {actual}")

    files = [validate_one(OUTPUT_DIR / name) for name in EXPECTED if (OUTPUT_DIR / name).exists()]
    total_pages = sum(item["pages"] or 0 for item in files)
    narrative_pages_with_internal_registers = sum(item["pages"] or 0 for item in files[:5])
    total_trip_notes = sum(item["trip_report_source_notes"] for item in files)
    all_errors = top_errors + [f"{item['file']}: {error}" for item in files for error in item["errors"]]
    if files and files[0]["pages"] and files[0]["pages"] > 3:
        all_errors.append("Cover letter exceeds the three-page maximum")
    if narrative_pages_with_internal_registers > 40:
        all_errors.append(
            f"Sections 8.4.1-8.4.5 total {narrative_pages_with_internal_registers} pages, exceeding 40"
        )
    report = {
        "output_directory": str(OUTPUT_DIR),
        "expected_file_count": 7,
        "actual_file_count": len(actual),
        "total_pages": total_pages,
        "narrative_pages_8_4_1_through_8_4_5_including_internal_registers": narrative_pages_with_internal_registers,
        "total_trip_report_source_notes": total_trip_notes,
        "files": files,
        "errors": all_errors,
        "status": "PASS" if not all_errors else "FAIL",
    }
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(json.dumps(report, indent=2), encoding="utf-8")

    print(f"Status: {report['status']}")
    print(f"Files: {len(actual)} | Pages: {total_pages} | Trip-report source notes: {total_trip_notes}")
    for item in files:
        print(
            f"{item['file']}: {item['pages']} pages, {item['words']} words, "
            f"{item['trip_report_source_notes']} trip notes, "
            f"{len(item['placeholders'])} placeholder types, "
            f"{len(item['errors'])} errors, {len(item['warnings'])} warnings"
        )
    if all_errors:
        print("Errors:")
        for error in all_errors:
            print(f"- {error}")
    return 0 if not all_errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
