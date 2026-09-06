#!/usr/bin/env python3
"""
validate_v2.py — QC checks for a Jacobs-branded proposal draft / rendered docx.

Usage:
    python work/validate_v2.py <draft.md> [rendered.docx] --spec <spec.md|.json> \
        --registry <registry.json> --budget-pages N [--required-heads file] [--json out]

Runs 8 checks, each reported PASS / FAIL / SKIP. The script itself always
exits 0 (per-check results are the signal, not the process exit code) unless
argument parsing fails.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

sys.path.insert(0, str(Path(__file__).resolve().parent))
try:
    from build_docx import TAG_RE, split_draft_notes  # reuse the same tag regex / splitter
except Exception:
    TAG_RE = re.compile(r"\[(?:PLACEHOLDER|VERIFY|CONFIRM|TBD)[^\]]*\]")

    def split_draft_notes(md_text: str):
        lines = md_text.splitlines()
        for idx, line in enumerate(lines):
            if line.strip().lower() == "## draft notes":
                return "\n".join(lines[:idx]), "\n".join(lines[idx:])
        return md_text, None

ALLOWED_HEX = {
    "231EDC", "0A7DFF", "5AE6FF", "001E55", "000000", "FFFFFF",
    "333333", "A5A5A5", "C8C8C8", "E6E6E6", "FFFF00",
}

LEAKAGE_TERMS = ["[CLIENT]", "Hull", "Town of Hull", "Santa Monica", "SWIP", "Helia", "Riverbend"]

DEFAULT_BANNED_WORDS = [
    "world-class", "best-in-class", "leverage", "synergy", "cutting-edge",
    "state-of-the-art", "robust", "seamless", "holistic", "utilize",
    "passionate", "unparalleled", "proven track record", "we are pleased",
    "as a leading",
]

EXCL_RE = re.compile(
    r"§\s?\d+(?:\.\d+)*"
    r"|\bTable\s+\d+[-.]\d+\b"
    r"|\bFigure\s+\d+[-.]\d+\b"
    r"|\b(?:page|pp?\.)\s*\d+\b",
    re.IGNORECASE,
)

NUM_RE = re.compile(
    r"(?P<money>\$\s?\d[\d,]*(?:\.\d+)?\s?(?:million|billion|M|B)?)"
    r"|(?P<pct>\d+(?:\.\d+)?%)"
    r"|(?P<unit>\d+(?:\.\d+)?\s?(?:MGD|mgd|miles|mi)\b)"
    r"|(?P<dec>\d+\.\d+)"
    r"|(?P<int>\b\d{2,}\b)"
)

TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------

def bare_number(token: str) -> str:
    m = re.search(r"\d+(?:\.\d+)?", token.replace(",", ""))
    return m.group(0) if m else token


def strip_trailing_zero(num: str) -> str:
    if "." in num:
        num = num.rstrip("0").rstrip(".")
    return num


def normalize(num: str) -> str:
    return strip_trailing_zero(bare_number(num))


def walk_json_numbers(obj, out: set) -> None:
    if isinstance(obj, dict):
        for v in obj.values():
            walk_json_numbers(v, out)
    elif isinstance(obj, list):
        for v in obj:
            walk_json_numbers(v, out)
    elif isinstance(obj, (int, float)) and not isinstance(obj, bool):
        out.add(normalize(str(obj)))
    elif isinstance(obj, str):
        for m in re.finditer(r"\d+(?:\.\d+)?", obj):
            out.add(normalize(m.group(0)))


def load_allowed_numbers(spec_path: Path | None, registry_path: Path | None) -> set:
    allowed = set()
    if spec_path and spec_path.exists():
        if spec_path.suffix.lower() == ".json":
            try:
                data = json.loads(spec_path.read_text(encoding="utf-8"))
                walk_json_numbers(data, allowed)
            except Exception:
                pass
        else:
            text = spec_path.read_text(encoding="utf-8")
            lines = text.splitlines()
            # "Locked number sheet" section (until next heading of <= its level)
            for idx, line in enumerate(lines):
                m = re.match(r"^(#{1,6})\s*Locked number sheet\s*$", line.strip(), re.IGNORECASE)
                if m:
                    level = len(m.group(1))
                    j = idx + 1
                    chunk = []
                    while j < len(lines):
                        hm = re.match(r"^(#{1,6})\s", lines[j])
                        if hm and len(hm.group(1)) <= level:
                            break
                        chunk.append(lines[j])
                        j += 1
                    for n in re.finditer(r"\d+(?:\.\d+)?", "\n".join(chunk)):
                        allowed.add(normalize(n.group(0)))
                    break
            # Any table with a "Value" column
            for idx, line in enumerate(lines):
                if line.strip().startswith("|") and "value" in line.lower():
                    cols = [c.strip().lower() for c in line.strip().strip("|").split("|")]
                    if "value" in cols:
                        val_idx = cols.index("value")
                        j = idx + 1
                        if j < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[j]):
                            j += 1
                        while j < len(lines) and lines[j].strip().startswith("|"):
                            cells = [c.strip() for c in lines[j].strip().strip("|").split("|")]
                            if val_idx < len(cells):
                                for n in re.finditer(r"\d+(?:\.\d+)?", cells[val_idx]):
                                    allowed.add(normalize(n.group(0)))
                            j += 1
    if registry_path and registry_path.exists():
        try:
            data = json.loads(registry_path.read_text(encoding="utf-8"))
            walk_json_numbers(data, allowed)
        except Exception:
            pass
    return allowed


def get_body_text_md(draft_md: str) -> str:
    main_md, _ = split_draft_notes(draft_md)
    return main_md


def get_docx_text(docx_path: Path) -> tuple[str, list]:
    """Returns (full_text, list_of_paragraph_texts) covering body paragraphs + table cells."""
    try:
        from docx import Document
    except Exception:
        return "", []
    doc = Document(str(docx_path))
    texts = []

    def walk_block_items(container):
        for p in container.paragraphs:
            texts.append(p.text)
        for t in container.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        texts.append(p.text)
    walk_block_items(doc)
    return "\n".join(texts), texts


def result(name: str, status: str, detail: str = "") -> dict:
    return {"check": name, "status": status, "detail": detail}


def print_result(r: dict) -> None:
    line = f"[{r['status']}] {r['check']}"
    if r["detail"]:
        line += f" - {r['detail']}"
    print(line)


# --------------------------------------------------------------------------
# Checks
# --------------------------------------------------------------------------

def check_tags(draft_md: str, docx_path: Path | None) -> dict:
    if docx_path:
        text, _ = get_docx_text(docx_path)
    else:
        text = get_body_text_md(draft_md)
    hits = TAG_RE.findall(text)
    if hits:
        return result("1. No open tags in body", "FAIL", f"{len(hits)} found: {hits[:5]}")
    return result("1. No open tags in body", "PASS")


HEADING_NUM_RE = re.compile(r"^#{1,6}\s+(\d+(?:\.\d+)*)\b")


def check_numbers(draft_md: str, spec_path: Path | None, registry_path: Path | None) -> dict:
    body = get_body_text_md(draft_md)
    allowed = load_allowed_numbers(spec_path, registry_path)
    excl_spans = [(m.start(), m.end()) for m in EXCL_RE.finditer(body)]
    # Section/subsection numbers in headings ("## 3.1 Heading") are
    # structural, not facts that need to trace to a source.
    pos = 0
    for line in body.splitlines(keepends=True):
        hm = HEADING_NUM_RE.match(line)
        if hm:
            excl_spans.append((pos + hm.start(1), pos + hm.end(1)))
        pos += len(line)

    def in_excl(pos):
        return any(s <= pos < e for s, e in excl_spans)

    unmatched = []
    for m in NUM_RE.finditer(body):
        if in_excl(m.start()):
            continue
        token = m.group(0)
        if m.lastgroup == "int" and len(token) == 4 and token.isdigit() and 1900 <= int(token) <= 2100:
            continue
        if normalize(token) in allowed:
            continue
        unmatched.append(token)
    if unmatched:
        uniq = sorted(set(unmatched))
        return result("2. Numbers traceable to spec/registry", "FAIL",
                       f"{len(uniq)} unmatched: {uniq[:15]}")
    return result("2. Numbers traceable to spec/registry", "PASS")


def check_pages(draft_md: str, docx_path: Path | None, budget: int | None) -> tuple[dict, int | None]:
    if not docx_path or budget is None:
        return result("3. Page budget", "SKIP", "no rendered docx or --budget-pages given"), None
    try:
        import win32com.client as win32
    except Exception:
        return result("3. Page budget", "SKIP", "pywin32/Word COM unavailable"), None
    pages = None
    try:
        word = win32.gencache.EnsureDispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(docx_path.resolve()))
        try:
            pages = doc.ComputeStatistics(2)  # wdStatisticPages
        except Exception:
            pages = None
        if pages is None:
            try:
                import fitz  # PyMuPDF
                tmp_pdf = docx_path.with_suffix(".validate_tmp.pdf")
                doc.ExportAsFixedFormat(str(tmp_pdf.resolve()), 17)
                pdf = fitz.open(str(tmp_pdf))
                pages = pdf.page_count
                pdf.close()
                tmp_pdf.unlink(missing_ok=True)
            except Exception:
                pages = None
        doc.Close(False)
        word.Quit()
    except Exception:
        return result("3. Page budget", "SKIP", "Word COM call failed"), None
    if pages is None:
        return result("3. Page budget", "SKIP", "could not determine page count"), None
    status = "PASS" if pages <= budget else "FAIL"
    return result("3. Page budget", status, f"{pages} pages (budget {budget})"), pages


def count_devices_md(draft_md: str) -> int:
    body = get_body_text_md(draft_md)
    lines = body.splitlines()
    n = 0
    n += sum(1 for l in lines if re.match(r"^\s*>\s*STAT:", l))
    n += sum(1 for l in lines if re.match(r"^\s*>\s*QUOTE:", l))
    n += sum(1 for l in lines if re.match(r"^\s*>\s*FACTBOX:", l))
    n += sum(1 for l in lines if re.match(r"^\s*>\s*CALLOUT:", l))
    n += sum(1 for l in lines if TABLE_SEP_RE.match(l))  # one separator row per GFM table
    return n


def check_devices(draft_md: str, pages: int | None) -> dict:
    if pages is None:
        return result("4. Device density (>= pages-1)", "SKIP", "page count unavailable")
    devices = count_devices_md(draft_md)
    threshold = max(pages - 1, 0)
    status = "PASS" if devices >= threshold else "FAIL"
    return result("4. Device density (>= pages-1)", status,
                   f"{devices} devices, need >= {threshold}")


def find_hex_colors(el) -> set:
    found = set()
    for e in el.iter():
        for attr in ("val", "color", "fill"):
            v = e.get(W_NS + attr)
            if v and re.fullmatch(r"[0-9A-Fa-f]{6}", v):
                found.add(v.upper())
    return found


def check_colors(docx_path: Path | None) -> dict:
    """Scans document.xml fully, and styles.xml only for styles actually
    referenced from document.xml -- python-docx's default template ships a
    full built-in style gallery (Heading 3-9, table styles, etc.) whose
    hardcoded theme swatches are never applied to visible content, and
    would otherwise false-positive on every generated document."""
    if not docx_path:
        return result("5. Allowed brand colors only", "SKIP", "no rendered docx given")
    try:
        with zipfile.ZipFile(docx_path) as z:
            doc_xml = z.read("word/document.xml")
            styles_xml = z.read("word/styles.xml") if "word/styles.xml" in z.namelist() else None
    except Exception as e:
        return result("5. Allowed brand colors only", "SKIP", f"could not read docx: {e}")

    bad = set()
    try:
        doc_root = ET.fromstring(doc_xml)
    except Exception as e:
        return result("5. Allowed brand colors only", "SKIP", f"could not parse document.xml: {e}")
    bad |= find_hex_colors(doc_root) - ALLOWED_HEX

    if styles_xml is not None:
        used_ids = {"Normal"}
        for tag in ("pStyle", "rStyle", "tblStyle"):
            for el in doc_root.iter(W_NS + tag):
                v = el.get(W_NS + "val")
                if v:
                    used_ids.add(v)
        try:
            styles_root = ET.fromstring(styles_xml)
            for dd in styles_root.findall(W_NS + "docDefaults"):
                bad |= find_hex_colors(dd) - ALLOWED_HEX
            for style_el in styles_root.findall(W_NS + "style"):
                sid = style_el.get(W_NS + "styleId")
                if sid in used_ids:
                    bad |= find_hex_colors(style_el) - ALLOWED_HEX
        except Exception:
            pass

    if bad:
        return result("5. Allowed brand colors only", "FAIL", f"disallowed hex: {sorted(bad)}")
    return result("5. Allowed brand colors only", "PASS")


def check_leakage(draft_md: str, docx_path: Path | None) -> dict:
    body = get_body_text_md(draft_md)
    if docx_path:
        docx_text, _ = get_docx_text(docx_path)
        body = body + "\n" + docx_text
    found = [term for term in LEAKAGE_TERMS if term in body]
    if found:
        return result("6. No cross-client leakage", "FAIL", f"found: {found}")
    return result("6. No cross-client leakage", "PASS")


def load_banned_words() -> list:
    guide = Path(r"C:\Users\micha\Desktop\Wiki\voice\voice-guide.md")
    if guide.exists():
        text = guide.read_text(encoding="utf-8")
        m = re.search(r"^#{1,6}\s*Banned words\s*$", text, re.IGNORECASE | re.MULTILINE)
        if m:
            rest = text[m.end():]
            nxt = re.search(r"^#{1,6}\s", rest, re.MULTILINE)
            chunk = rest[:nxt.start()] if nxt else rest
            words = re.findall(r"^[-*]\s*(.+)$", chunk, re.MULTILINE)
            words = [w.strip().strip("`").split(" (")[0].strip() for w in words if w.strip()]
            if words:
                return words
    return DEFAULT_BANNED_WORDS


def check_banned_words(draft_md: str, docx_path: Path | None) -> dict:
    body = get_body_text_md(draft_md)
    if docx_path:
        docx_text, _ = get_docx_text(docx_path)
        body = body + "\n" + docx_text
    body_l = body.lower()
    banned = load_banned_words()
    found = [w for w in banned if w.lower() in body_l]
    if found:
        return result("7. No banned voice words", "FAIL", f"found: {found}")
    return result("7. No banned voice words", "PASS")


def check_required_heads(draft_md: str, required_heads_path: Path | None) -> dict:
    if not required_heads_path:
        return result("8. Required headings present", "SKIP", "no --required-heads given")
    required = [l.strip() for l in required_heads_path.read_text(encoding="utf-8").splitlines() if l.strip()]
    headings = [l.strip("# ").strip() for l in draft_md.splitlines() if l.strip().startswith("#")]
    headings_l = [h.lower() for h in headings]
    missing = [r for r in required if r.lower() not in headings_l and
               not any(r.lower() in h for h in headings_l)]
    if missing:
        return result("8. Required headings present", "FAIL", f"missing: {missing}")
    return result("8. Required headings present", "PASS")


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(description="QC checks for a Jacobs proposal draft / rendered docx")
    parser.add_argument("draft_md", type=Path)
    parser.add_argument("rendered_docx", type=Path, nargs="?", default=None)
    parser.add_argument("--spec", type=Path, default=None)
    parser.add_argument("--registry", type=Path, default=None)
    parser.add_argument("--budget-pages", type=int, default=None)
    parser.add_argument("--required-heads", type=Path, default=None)
    parser.add_argument("--json", type=Path, default=None)
    args = parser.parse_args()

    draft_md = args.draft_md.read_text(encoding="utf-8")
    docx_path = args.rendered_docx if args.rendered_docx and args.rendered_docx.exists() else None

    results = []
    results.append(check_tags(draft_md, docx_path))
    results.append(check_numbers(draft_md, args.spec, args.registry))
    pages_result, pages = check_pages(draft_md, docx_path, args.budget_pages)
    results.append(pages_result)
    results.append(check_devices(draft_md, pages))
    results.append(check_colors(docx_path))
    results.append(check_leakage(draft_md, docx_path))
    results.append(check_banned_words(draft_md, docx_path))
    results.append(check_required_heads(draft_md, args.required_heads))

    for r in results:
        print_result(r)

    n_fail = sum(1 for r in results if r["status"] == "FAIL")
    n_pass = sum(1 for r in results if r["status"] == "PASS")
    n_skip = sum(1 for r in results if r["status"] == "SKIP")
    print(f"\n{n_pass} PASS, {n_fail} FAIL, {n_skip} SKIP")

    if args.json:
        args.json.write_text(json.dumps(results, indent=2), encoding="utf-8")

    return 0


if __name__ == "__main__":
    sys.exit(main())
